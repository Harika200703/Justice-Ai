from pathlib import Path
from pypdf import PdfReader
from docx import Document
SUPPORTED = {".pdf", ".docx", ".txt"}
def extract_text(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf": text = "\n\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    elif ext == ".docx":
        doc = Document(path); parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows: parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext == ".txt": text = Path(path).read_text(encoding="utf-8", errors="replace")
    else: raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")
    if not text.strip(): raise ValueError("No selectable text was extracted. Scanned PDFs need OCR support.")
    return text.strip()
